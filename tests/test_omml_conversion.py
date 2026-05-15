#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from omml_to_latex import OmmlConversionStats, convert_omml_xml  # noqa: E402

NS = {
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}


def omath(inner: str) -> str:
    return f'<m:oMath xmlns:m="{NS["m"]}" xmlns:w="{NS["w"]}">{inner}</m:oMath>'


def r(text: str) -> str:
    return f"<m:r><m:t>{text}</m:t></m:r>"


def insert_omath(paragraph, xml: str) -> None:
    paragraph._p.append(etree.fromstring(xml.encode("utf-8")))


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("1 测试公式")

    p = doc.add_paragraph("行内公式 ")
    insert_omath(
        p,
        omath(
            "<m:f>"
            f"<m:num><m:r><m:t>a+b</m:t></m:r></m:num>"
            f"<m:den><m:r><m:t>c</m:t></m:r></m:den>"
            "</m:f>"
        ),
    )
    p.add_run(" 结束")

    p = doc.add_paragraph()
    p._p.append(
        etree.fromstring(
            f'''<m:oMathPara xmlns:m="{NS['m']}" xmlns:w="{NS['w']}">
              <m:oMath><m:sSup><m:e>{r('x')}</m:e><m:sup>{r('2')}</m:sup></m:sSup></m:oMath>
            </m:oMathPara>'''.encode("utf-8")
        )
    )

    p = doc.add_paragraph()
    p._p.append(
        etree.fromstring(
            f'''<m:oMathPara xmlns:m="{NS['m']}" xmlns:w="{NS['w']}">
              <m:oMath><m:f><m:num>{r('E')}</m:num><m:den>{r('mc')}</m:den></m:f></m:oMath>
            </m:oMathPara>'''.encode("utf-8")
        )
    )
    p.add_run("（2-1）")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "变量"
    table.cell(0, 1).text = "公式"
    table.cell(1, 0).text = "平方根"
    cell_p = table.cell(1, 1).paragraphs[0]
    insert_omath(cell_p, omath(f"<m:rad><m:e>{r('x')}</m:e></m:rad>"))

    doc.save(path)


class OmmlConversionTest(unittest.TestCase):
    def test_core_omml_to_latex(self) -> None:
        cases = [
            (
                "subsup",
                omath(
                    "<m:sSubSup>"
                    f"<m:e>{r('x')}</m:e><m:sub>{r('i')}</m:sub><m:sup>{r('2')}</m:sup>"
                    "</m:sSubSup>"
                ),
                "{x}_{i}^{2}",
            ),
            (
                "sum",
                omath(
                    '<m:nary><m:naryPr><m:chr m:val="∑"/></m:naryPr>'
                    f"<m:sub>{r('i=1')}</m:sub><m:sup>{r('n')}</m:sup>"
                    f"<m:e><m:sSub><m:e>{r('x')}</m:e><m:sub>{r('i')}</m:sub></m:sSub></m:e>"
                    "</m:nary>"
                ),
                r"\sum_{i = 1}^{n} {x}_{i}",
            ),
            (
                "delimiter",
                omath(
                    '<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
                    f"<m:e>{r('a+b')}</m:e></m:d>"
                ),
                r"\left(a + b\right)",
            ),
            (
                "bar_without_props",
                omath(f"<m:bar><m:e>{r('x')}</m:e></m:bar>"),
                r"\overline{x}",
            ),
        ]
        for name, xml, expected in cases:
            with self.subTest(name=name):
                stats = OmmlConversionStats()
                tex = convert_omml_xml(xml, stats)
                self.assertEqual(tex, expected)
                self.assertEqual(stats.detected, 1)
                self.assertEqual(stats.converted, 1)
                self.assertFalse(stats.errors)

    def test_word_to_standard_md_preserves_inline_display_and_table_math(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            td_path = Path(td)
            docx = td_path / "math.docx"
            out = td_path / "math.standard.md"
            make_docx(docx)
            proc = subprocess.run(
                [sys.executable, str(SCRIPTS / "word_to_standard_md.py"), str(docx), "-o", str(out)],
                text=True,
                capture_output=True,
                check=True,
            )
            self.assertIn("已转换：4", proc.stdout)
            md = out.read_text(encoding="utf-8")
            self.assertIn(r"$\frac{a + b}{c}$", md)
            self.assertIn("::: equation", md)
            self.assertIn(r"{x}^{2}", md)
            self.assertIn("tag: 2-1", md)
            self.assertIn(r"\frac{E}{mc}", md)
            self.assertIn(r"$\sqrt{x}$", md)
            report = json.loads(out.with_suffix(".report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["omml_count"], 4)
            self.assertEqual(report["omml_converted"], 4)
            self.assertEqual(report["omml_unconverted"], 0)
            self.assertFalse(report["omml_errors"])


if __name__ == "__main__":
    unittest.main()
